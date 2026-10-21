#!/usr/bin/env python3
"""Выгрузка карточек из historias.xlsx в файлы YAML (по одному на событие).

Языки: pt, en, ru. Русский добавлен как третий язык: в YAML у каждого поля
{pt, en} теперь есть ключ ru.

Использование:
    python export_yaml.py historias.xlsx путь/к/src/content/events

Нужны: Python 3, пакеты openpyxl и pyyaml (pip install openpyxl pyyaml).
Скрипт читает листы «Карточки», «Тексты», «Места», «Источники», «Картинки»,
связывая их по id. Выгружаются все строки «Карточек»; в YAML статус будет
ready только если «Статус работы» = «готово», иначе draft.
Предупреждения печатаются в консоль.
"""
import json
import os
import sys

import yaml
from openpyxl import load_workbook

LANGS = ("pt", "en", "ru")


class Folded(str):
    pass


class Literal(str):
    pass


def _folded(dumper, data):
    return dumper.represent_scalar("tag:yaml.org,2002:str", str(data), style=">")


def _literal(dumper, data):
    return dumper.represent_scalar("tag:yaml.org,2002:str", str(data), style="|")


yaml.add_representer(Folded, _folded)
yaml.add_representer(Literal, _literal)


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def num(v):
    if v is None or str(v).strip() == "":
        return None
    f = float(str(v).replace(",", "."))
    return int(f) if f.is_integer() else f


def rows(ws):
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {h: v for h, v in zip(head, r) if h}
        if any(v is not None and str(v).strip() != "" for v in d.values()):
            yield d


def long(s):
    if s and "\n" in s:                      # абзацы сохраняются
        return Literal(s.replace("\r\n", "\n").strip())
    return Folded(s) if s and len(s) > 60 else s


def multi(d, keys, fold=True, required=True):
    """Собирает {pt, en, ru} из ячеек. keys: ключи заголовков в порядке LANGS.
    Если нет ни одного языка: None (или TODO для обязательных). Недостающий язык: TODO."""
    vals = [txt(d.get(k)) for k in keys]
    if all(v is None for v in vals):
        return {lg: "TODO" for lg in LANGS} if required else None
    return {lg: (long(v) if (fold and v) else (v or "TODO")) for lg, v in zip(LANGS, vals)}


def main(xlsx, outdir):
    wb = load_workbook(xlsx, data_only=True)
    texts = {r["id"]: r for r in rows(wb["Тексты"]) if r.get("id")}
    places = {r["id"]: r for r in rows(wb["Места"]) if r.get("id")}
    places_raw = {row[0]: row for row in wb["Места"].iter_rows(min_row=2, values_only=True) if row[0]}
    images = {r["id карточки"]: r for r in rows(wb["Картинки"]) if r.get("id карточки")}
    sources = {}
    for r in rows(wb["Источники"]):
        if r.get("Статус") == "нашли":
            sources.setdefault(r["id карточки"], []).append(r)

    os.makedirs(outdir, exist_ok=True)
    count = 0
    for c in rows(wb["Карточки"]):
        cid = txt(c.get("id"))
        if not cid:
            continue
        warn = []
        t, p = texts.get(cid), places.get(cid)
        if t is None:
            warn.append("нет строки в «Тексты»")
            t = {}
        if p is None:
            warn.append("нет строки в «Места»")
            p = {}

        rec = {"id": cid}
        pt_slug, en_slug, ru_slug = txt(p.get("Slug PT")), txt(p.get("Slug EN")), txt(p.get("Slug RU (латиницей)"))
        rec["slug"] = {"pt": pt_slug or "TODO", "en": en_slug or "TODO", "ru": ru_slug or "TODO"}
        rec["title"] = {"pt": txt(p.get("Título PT")) or "TODO", "en": txt(p.get("Title EN")) or "TODO",
                        "ru": txt(p.get("Title RU")) or "TODO"}
        rec["epoch"] = txt(c.get("Эпоха"))
        rec["level"] = txt(c.get("Уровень"))
        rec["year"] = num(c.get("Год"))
        yl = multi(p, ("Подпись года PT", "Подпись года EN", "Подпись года RU"), fold=False, required=False)
        if yl:
            rec["year_label"] = yl

        raw = places_raw.get(cid)
        if raw is not None:
            rec["place"] = {"name": txt(raw[7]), "lat": num(raw[8]) or 0.0, "lon": num(raw[9]) or 0.0}
            rec["bridge"] = {
                "name": txt(raw[10]) or "TODO",
                "address": txt(raw[11]) or "TODO",
                "lat": num(raw[12]) or 0.0,
                "lon": num(raw[13]) or 0.0,
                "url": txt(raw[14]) or "TODO",
            }
            if str(raw[15] or "").strip() != "да":
                warn.append("координаты не проверены")
        else:
            rec["place"] = {"name": "TODO", "lat": 0.0, "lon": 0.0}
            rec["bridge"] = {"name": "TODO", "address": "TODO", "lat": 0.0, "lon": 0.0, "url": "TODO"}
        rec["threads"] = [x for x in (txt(c.get("Нить 1")), txt(c.get("Нить 2"))) if x]
        rec["importance"] = int(num(c.get("Важность 1–3")) or 0)
        rec["label"] = txt(c.get("Метка"))
        rec["min_age"] = str(c.get("Возраст")).strip() if c.get("Возраст") is not None else None
        rec["sensitive"] = txt(c.get("Тяжёлая тема")) == "да"
        rec["hook"] = multi(t, ("Крючок PT (до 12 сл.)", "Крючок EN (до 12 сл.)", "Крючок RU (до 12 сл.)"))
        rec["story"] = multi(t, ("История PT (40–90 сл.)", "История EN (40–90 сл.)", "История RU (40–90 сл.)"))
        rec["more"] = multi(t, ("Подробнее PT (до 250 сл.)", "Подробнее EN (до 250 сл.)", "Подробнее RU (до 250 сл.)"))
        rec["sources"] = [
            {"title": s["Название (автор, заглавие, страница)"], "url": txt(s.get("URL")) or "TODO",
             "type": txt(s.get("Тип")), "accessed": str(s.get("Дата доступа") or "")[:10]}
            for s in sources.get(cid, [])
        ]
        img = images.get(cid)
        if img and txt(img.get("Файл (src)")):
            rec["image"] = {
                "src": txt(img["Файл (src)"]),
                "alt": multi(img, ("Alt PT", "Alt EN", "Alt RU"), fold=False),
                "credit": txt(img.get("Автор")) or "TODO",
                "license": txt(img.get("Лицензия")) or "TODO",
            }
            cap = multi(img, ("Подпись PT", "Подпись EN", "Подпись RU"), required=False)
            if cap:
                rec["image"]["caption"] = cap
        rel = txt(c.get("Связанные id (через запятую)"))
        rec["related"] = [x.strip() for x in rel.split(",") if x.strip()] if rel else []
        itype = txt(c.get("Интерактив: тип"))
        if itype:
            try:
                data = json.loads(c.get("Интерактив: данные (JSON)") or "{}")
            except json.JSONDecodeError as e:
                warn.append(f"JSON интерактива не читается: {e}")
                data = {}
            rec["interactive"] = {"type": itype, "data": data}

        ready = txt(c.get("Статус работы")) == "готово"
        if ready:
            if txt(c.get("Рецензент")):
                rec["reviewed_by"] = txt(c.get("Рецензент"))
            if c.get("Дата рецензии"):
                rec["reviewed_on"] = str(c.get("Дата рецензии"))[:10]
            else:
                warn.append("готово, но нет даты рецензии")
        rec["status"] = "ready" if ready else "draft"
        rec["placeholder"] = False

        for key in ("hook", "story", "more"):
            if any(v == "TODO" for v in rec[key].values()):
                warn.append(f"в {key} не хватает языка")

        path = os.path.join(outdir, f"{cid}.yaml")
        with open(path, "w", encoding="utf-8") as fh:
            yaml.dump(rec, fh, allow_unicode=True, sort_keys=False, width=78, default_flow_style=False)
        count += 1
        print(f"{cid}: записан" + (" | ⚠ " + "; ".join(warn) if warn else ""))
    print(f"Готово: {count} файлов в {outdir}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
